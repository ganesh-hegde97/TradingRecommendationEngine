from openpyxl import load_workbook

from python_app.exports.excel_export import export_workbook
from python_app.services.recommendation_service import RecommendationService


def test_export_accepts_typed_recommendations(tmp_path) -> None:
    output = tmp_path / "recommendations.xlsx"
    export_workbook(RecommendationService().screen_demo(), output, "test data")
    workbook = load_workbook(output, data_only=False)
    assert output.exists()
    assert workbook.sheetnames == [
        "Recommendations",
        "All Candidates",
        "Parameters",
        "Source Data",
        "Checks",
    ]
    assert workbook["Source Data"]["A5"].value
